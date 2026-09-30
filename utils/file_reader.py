from io import BytesIO

from docx import Document
from pypdf import PdfReader


def _read_pdf(file_bytes: bytes) -> str:
    reader = PdfReader(BytesIO(file_bytes))
    pages = []

    for page in reader.pages:
        text = page.extract_text()
        if text:
            pages.append(text)

    return "\n".join(pages)


def _read_docx(file_bytes: bytes) -> str:
    document = Document(BytesIO(file_bytes))
    parts = []

    for paragraph in document.paragraphs:
        if paragraph.text.strip():
            parts.append(paragraph.text)

    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    parts.append(cell.text)

    return "\n".join(parts)


def extract_text_from_uploaded_file(uploaded_file) -> str:
    file_name = uploaded_file.name.lower()
    file_bytes = uploaded_file.getvalue()

    if file_name.endswith(".pdf"):
        return _read_pdf(file_bytes)

    if file_name.endswith(".docx"):
        return _read_docx(file_bytes)

    raise ValueError("Only PDF and DOCX resumes are supported.")
